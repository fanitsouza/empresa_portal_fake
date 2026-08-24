"""
Processo 3 - Cadastro Automatizado no Portal Fake / ERP
Portal Fake Soluções Digitais - Hyperautomation

Responsabilidade:
1. Carregar registros aprovados da Planilha Mestra (.xlsx).
2. Abrir o Portal Fake e consultar o CPF do cliente.
3. Se o cliente JÁ estiver cadastrado:
   - Registra como cadastro duplicado.
   - Atualiza o status na planilha mestra.
   - Salva evidência e log na pasta 'Cadastro_com_erro'.
4. Se o cliente NÃO estiver cadastrado:
   - Preenche o formulário e salva o cadastro.
   - Valida se o cadastro foi realizado com sucesso:
     * Se falhou: registra erro, atualiza a planilha e salva em 'Cadastro_com_erro'.
     * Se sucesso: atualiza o status na planilha e salva em 'Cadastro_Concluido'.
5. Retorna os dados consolidados para o Processo 4 (SAC).
"""

from __future__ import annotations

import argparse
from datetime import datetime
import json
import logging
import os
from pathlib import Path
import re
from typing import Any, Dict, List, Tuple

from dotenv import load_dotenv
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

BASE_DIR = Path(__file__).resolve().parent
PATH_ROOT = BASE_DIR.parent
RESOURCES_DIR = PATH_ROOT / "resources"

INDEX_HTML = RESOURCES_DIR / "portal_fake" / "index.html"
PLANILHA_PADRAO = RESOURCES_DIR / "planilha_mestra.xlsx"
BROWSER_DATA_DIR = RESOURCES_DIR / "browser_data"
PASTA_SCREENSHOTS = RESOURCES_DIR / "screenshots"
PASTA_LOGS = RESOURCES_DIR / "logs"

# Pastas de saída das regras de negócio do Processo 3
PASTA_ERP = RESOURCES_DIR / "ERP_Portal_Fake"
PASTA_CADASTRO_ERRO = PASTA_ERP / "Cadastro_com_erro"
PASTA_CADASTRO_CONCLUIDO = PASTA_ERP / "Cadastro_Concluido"

load_dotenv(BASE_DIR / ".env")

import sys
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from planilha_mestra import atualizar_status_cadastro_planilha

LOGGER = logging.getLogger("PROCESSO_3_CADASTRO")
if not LOGGER.handlers:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | [%(name)s] %(levelname)s: %(message)s")


def garantir_diretorios() -> None:
    """Cria os diretórios necessários para evidências e logs."""
    for pasta in (
        PASTA_CADASTRO_ERRO,
        PASTA_CADASTRO_CONCLUIDO,
        PASTA_SCREENSHOTS,
        PASTA_LOGS,
    ):
        pasta.mkdir(parents=True, exist_ok=True)


def carregar_clientes_planilha(caminho_planilha: str | Path) -> List[Dict[str, Any]]:
    """Carrega os registros com status de extração válido da Planilha Mestra."""
    caminho = Path(caminho_planilha).expanduser().resolve()
    if not caminho.is_file():
        raise FileNotFoundError(f"Planilha mestra não encontrada: {caminho}")

    workbook = load_workbook(caminho, data_only=True)
    try:
        nome_aba = "Dados extraidos"
        planilha = workbook[nome_aba] if nome_aba in workbook.sheetnames else workbook.active
        linhas = list(planilha.iter_rows(values_only=True))
        if not linhas:
            return []

        cabecalho = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(linhas[0])]
        col_indices = {nome: idx for idx, nome in enumerate(cabecalho)}

        obrigatorias = {"Nome", "CPF", "E-mail"}
        if not obrigatorias.issubset(set(col_indices.keys())):
            raise ValueError(f"Colunas obrigatórias ausentes na planilha: {obrigatorias - set(col_indices.keys())}")

        clientes: List[Dict[str, Any]] = []
        for idx, linha in enumerate(linhas[1:], start=2):
            if not any(linha):
                continue
            
            sucesso = str(linha[col_indices.get("Sucesso", 7)] or "Sim").strip().casefold()
            if sucesso not in {"sim", "true", "1", "sucesso"}:
                continue

            nome_completo = str(linha[col_indices["Nome"]] or "").strip()
            partes = nome_completo.split(" ", 1)
            primeiro_nome = partes[0]
            sobrenome = partes[1] if len(partes) > 1 else ""

            cpf_bruto = str(linha[col_indices["CPF"]] or "").strip()
            cpf_digitos = re.sub(r"\D", "", cpf_bruto)

            cliente = {
                "linha": idx,
                "nome_completo": nome_completo,
                "nome": primeiro_nome,
                "sobrenome": sobrenome,
                "cpf": cpf_digitos,
                "cpf_formatado": cpf_bruto,
                "email": str(linha[col_indices.get("E-mail", 3)] or "").strip(),
                "telefone": str(linha[col_indices.get("Telefone", 4)] or "").strip(),
                "nascimento": str(linha[col_indices.get("Nascimento", 5)] or "").strip(),
                "endereco": str(linha[col_indices.get("Endereco", 6)] or "").strip(),
                "status": "ATIVO",
                "observacao": "Cadastro realizado automaticamente via Robô RPA",
            }
            clientes.append(cliente)
        return clientes
    finally:
        workbook.close()


def consultar_cpf_no_portal(page, cpf: str) -> Tuple[bool, int]:
    """
    Consulta o CPF no Portal Fake utilizando o campo de busca rápida (#q).
    Retorna (ja_cadastrado, quantidade_encontrada).
    """
    cpf_digitos = re.sub(r"\D", "", cpf)
    page.fill("#q", cpf_digitos)
    page.click("#btnBuscar")
    page.wait_for_timeout(300)

    texto_count = page.locator("#count").inner_text().strip()
    qtd_linhas = page.locator("#tbody tr").count()

    ja_existe = qtd_linhas > 0 or ("0 encontrados" not in texto_count and "encontrado" in texto_count)

    # Limpa a busca para as próximas operações
    page.fill("#q", "")
    page.click("#btnLimpar")
    page.wait_for_timeout(200)

    return ja_existe, qtd_linhas


def cadastrar_cliente_portal(page, cliente: Dict[str, Any]) -> Tuple[bool, str]:
    """
    Preenche o formulário de cadastro no Portal Fake e valida o salvamento.
    Retorna (sucesso, mensagem_erro).
    """
    page.click("#btnNovo")
    page.locator("#modal").wait_for(state="visible", timeout=3000)

    page.fill("#f_nome", cliente.get("nome", ""))
    page.fill("#f_sobrenome", cliente.get("sobrenome", ""))
    page.fill("#f_cpf", cliente.get("cpf", ""))
    page.fill("#f_email", cliente.get("email", ""))
    page.fill("#f_telefone", cliente.get("telefone", ""))
    
    # Formatação de nascimento para input date (YYYY-MM-DD)
    nasc = cliente.get("nascimento", "")
    if "/" in nasc:
        partes = nasc.split("/")
        if len(partes) == 3:
            nasc = f"{partes[2]}-{partes[1].zfill(2)}-{partes[0].zfill(2)}"
    page.fill("#f_nascimento", nasc)

    page.fill("#f_endereco", cliente.get("endereco", ""))
    page.fill("#f_observacao", cliente.get("observacao", "Cadastro RPA"))
    page.select_option("#f_status", (cliente.get("status") or "ATIVO").upper())

    # Submete o formulário
    page.click("#btnSalvar")

    try:
        page.locator("#modal").wait_for(state="hidden", timeout=3000)
        return True, ""
    except Exception:
        # Se modal permanecer visível, captura a mensagem de erro do formulário
        msg_erro = page.locator("#formMsg").inner_text().strip() or "Erro na validação do formulário"
        try:
            page.click("#btnModalClose")
            page.locator("#modal").wait_for(state="hidden", timeout=2000)
        except Exception:
            pass
        return False, msg_erro


def salvar_evidencia(
    page,
    cliente: Dict[str, Any],
    tipo: str,
    motivo: str = "",
) -> Tuple[Path, Path]:
    """
    Salva uma captura de tela (screenshot) e um arquivo JSON descritivo na pasta adequada.
    tipo: 'erro' -> PASTA_CADASTRO_ERRO | 'concluido' -> PASTA_CADASTRO_CONCLUIDO
    """
    garantir_diretorios()
    pasta_destino = PASTA_CADASTRO_CONCLUIDO if tipo == "concluido" else PASTA_CADASTRO_ERRO
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    cpf_sufixo = cliente.get("cpf", "000")[-4:]

    nome_base = f"{tipo}_CPF_{cpf_sufixo}_{timestamp}"
    caminho_img = pasta_destino / f"{nome_base}.png"
    caminho_json = pasta_destino / f"{nome_base}.json"

    try:
        page.screenshot(path=str(caminho_img), full_page=True)
    except Exception as e:
        LOGGER.warning(f"Não foi possível salvar screenshot: {e}")

    dados_evidencia = {
        "tipo": tipo,
        "timestamp": datetime.now().isoformat(),
        "cliente": cliente,
        "motivo": motivo,
        "screenshot": caminho_img.name,
    }
    caminho_json.write_text(json.dumps(dados_evidencia, ensure_ascii=False, indent=2), encoding="utf-8")
    return caminho_img, caminho_json


def executar_processo_3(
    caminho_planilha: Path | str | None = None,
    quantidade: int = 0,
    headless: bool = False,
) -> Dict[str, Any]:
    """
    Executa o Processo 3 (Setor de Cadastro no Portal Fake):
    - Consulta prévia de CPF (detecção de duplicidade)
    - Cadastro automatizado via RPA Playwright
    - Validação de sucesso
    - Gravação nas pastas Cadastro_com_erro ou Cadastro_Concluido
    - Atualização da Planilha Mestra
    - Retorno estruturado para o Processo 4 (SAC)
    """
    print("\n" + "=" * 70)
    print("💻 [PROCESSO 3] INICIANDO SETOR DE CADASTRO NO PORTAL FAKE")
    print("=" * 70)

    garantir_diretorios()
    planilha = Path(caminho_planilha) if caminho_planilha else PLANILHA_PADRAO

    print(f"\n[Etapa 1] Lendo registros da Planilha Mestra: {planilha.name}")
    try:
        clientes = carregar_clientes_planilha(planilha)
    except Exception as e:
        LOGGER.error(f"Falha ao carregar a planilha mestra: {e}")
        return {"total": 0, "sucessos": 0, "duplicados": 0, "erros": 0, "resultados": []}

    if not clientes:
        print("ℹ️ Nenhum cliente apto para cadastro encontrado na planilha.")
        return {"total": 0, "sucessos": 0, "duplicados": 0, "erros": 0, "resultados": []}

    lista_processar = clientes[:quantidade] if quantidade > 0 else clientes
    total = len(lista_processar)
    print(f"  • {total} cliente(s) selecionado(s) para processamento.")

    resultados_detalhados: List[Dict[str, Any]] = []
    sucessos = 0
    duplicados = 0
    erros = 0

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(BROWSER_DATA_DIR),
            headless=headless,
        )
        page = context.new_page()
        page.on("dialog", lambda dialog: dialog.accept())

        url_portal = INDEX_HTML.resolve().as_uri()
        print(f"\n[Etapa 2] Acessando Portal Fake: {url_portal}")
        page.goto(url_portal)
        page.wait_for_selector("#tbody", timeout=5000)

        for i, cliente in enumerate(lista_processar, start=1):
            cpf_formatado = cliente.get("cpf_formatado", cliente.get("cpf"))
            nome = cliente.get("nome_completo")
            print(f"\n------------------------------------------------------------")
            print(f"👤 ({i}/{total}) Analisando: {nome} | CPF: {cpf_formatado}")

            # 1. Consulta CPF no Portal Fake
            print(f"  🔍 Consultando CPF {cpf_formatado} no Portal Fake...")
            ja_cadastrado, _ = consultar_cpf_no_portal(page, cliente["cpf"])

            if ja_cadastrado:
                motivo = "Cadastro duplicado: CPF já existente na base de dados do Portal Fake"
                print(f"  ⚠️ STATUS: CADASTRO DUPLICADO! ({motivo})")
                
                img_path, json_path = salvar_evidencia(page, cliente, tipo="erro", motivo=motivo)
                print(f"  📁 Evidência salva em: {PASTA_CADASTRO_ERRO.name}/{img_path.name}")
                
                duplicados += 1
                resultados_detalhados.append({
                    "cliente": cliente,
                    "cpf": cliente["cpf"],
                    "status_cadastro": "DUPLICADO",
                    "motivo_erro": motivo,
                    "data_cadastro": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                    "evidencia": str(img_path),
                })
                continue

            # 2. Preenche e salva novo cadastro
            print(f"  📝 CPF não cadastrado. Preenchendo formulário no portal...")
            cadastrado_ok, msg_falha = cadastrar_cliente_portal(page, cliente)

            if cadastrado_ok:
                print(f"  ✅ STATUS: CADASTRADO COM SUCESSO! (Ativo no Portal Fake)")
                img_path, json_path = salvar_evidencia(page, cliente, tipo="concluido", motivo="Sucesso no cadastro")
                print(f"  📁 Evidência salva em: {PASTA_CADASTRO_CONCLUIDO.name}/{img_path.name}")

                sucessos += 1
                resultados_detalhados.append({
                    "cliente": cliente,
                    "cpf": cliente["cpf"],
                    "status_cadastro": "SUCESSO",
                    "motivo_erro": "",
                    "data_cadastro": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                    "evidencia": str(img_path),
                })
            else:
                motivo = f"Erro no preenchimento/validação do formulário: {msg_falha}"
                print(f"  ❌ STATUS: ERRO NO CADASTRO! ({motivo})")
                img_path, json_path = salvar_evidencia(page, cliente, tipo="erro", motivo=motivo)
                print(f"  📁 Evidência salva em: {PASTA_CADASTRO_ERRO.name}/{img_path.name}")

                erros += 1
                resultados_detalhados.append({
                    "cliente": cliente,
                    "cpf": cliente["cpf"],
                    "status_cadastro": "ERRO_CADASTRO",
                    "motivo_erro": motivo,
                    "data_cadastro": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                    "evidencia": str(img_path),
                })

        context.close()

    # 3. Atualizar Planilha Mestra com status de auditoria
    print(f"\n[Etapa 3] Atualizando status de cadastro na Planilha Mestra...")
    try:
        atualizar_status_cadastro_planilha(resultados_detalhados, planilha)
        print(f"  📗 Planilha Mestra atualizada com sucesso: {planilha.name}")
    except Exception as e:
        LOGGER.warning(f"Não foi possível gravar status na planilha mestra: {e}")

    print("\n" + "=" * 70)
    print("📊 [PROCESSO 3] RESUMO DO SETOR DE CADASTRO:")
    print(f"  • Total analisados: {total}")
    print(f"  • Cadastros Concluídos (Sucesso): {sucessos}")
    print(f"  • Cadastros Duplicados: {duplicados}")
    print(f"  • Cadastros com Erro: {erros}")
    print(f"  • Pasta Concluídos: {PASTA_CADASTRO_CONCLUIDO}")
    print(f"  • Pasta Erros/Duplicados: {PASTA_CADASTRO_ERRO}")
    print("=" * 70 + "\n")

    return {
        "total": total,
        "sucessos": sucessos,
        "duplicados": duplicados,
        "erros": erros,
        "resultados": resultados_detalhados,
    }


def executar_processo_4(
    caminho_planilha: Path | str | None = None,
    quantidade: int = 0,
    headless: bool = False,
    **kwargs,
) -> Dict[str, Any]:
    """Alias de compatibilidade para executar_processo_3."""
    return executar_processo_3(
        caminho_planilha=caminho_planilha,
        quantidade=quantidade,
        headless=headless,
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Processo 3 - Setor de Cadastro Automatizado no Portal Fake com Consulta CPF e Validação"
    )
    parser.add_argument(
        "planilha",
        nargs="?",
        default=str(PLANILHA_PADRAO),
        help="Caminho da Planilha Mestra .xlsx (padrão: resources/planilha_mestra.xlsx).",
    )
    parser.add_argument(
        "--quantidade",
        type=int,
        default=0,
        help="Quantidade limite de cadastros a processar (0 para todos).",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Executa o navegador em modo oculto (sem janela visível).",
    )
    parser.add_argument(
        "--no-headless",
        action="store_true",
        help="Força exibição visual da janela do navegador.",
    )

    args = parser.parse_args()
    is_headless = args.headless and not args.no_headless

    executar_processo_3(
        caminho_planilha=args.planilha,
        quantidade=args.quantidade,
        headless=is_headless,
    )


if __name__ == "__main__":
    main()
