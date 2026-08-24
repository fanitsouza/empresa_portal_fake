from __future__ import annotations

import argparse
from datetime import date, datetime
import json
import logging
import os
from pathlib import Path
import re
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from dotenv import load_dotenv


PATH_ROOT = Path(__file__).resolve().parents[2]
INDEX_HTML = PATH_ROOT / "resources" / "portal_fake" / "index.html"
PLANILHA_PADRAO = PATH_ROOT / "resources" / "planilha_mestra.xlsx"
BROWSER_DATA_DIR = PATH_ROOT / "resources" / "browser_data"
PASTA_SCREENSHOTS = PATH_ROOT / "resources" / "screenshots"
PASTA_LOGS = PATH_ROOT / "resources" / "logs"
NOME_ABA = "Dados extraidos"
LOGGER = logging.getLogger("PROCESSO_3_CADASTRO")

load_dotenv(PATH_ROOT / "source" / ".env")


class IntegracaoAPIError(RuntimeError):
    """Falha controlada ao integrar com API ou serviço de ML."""


def configurar_logs(pasta_logs: str | Path = PASTA_LOGS) -> Path:
    """Configura logs no terminal e em arquivo persistente."""
    pasta = Path(pasta_logs).expanduser().resolve()
    pasta.mkdir(parents=True, exist_ok=True)
    caminho_log = pasta / "processo_3_cadastro.log"

    if not LOGGER.handlers:
        LOGGER.setLevel(logging.INFO)
        formato = logging.Formatter(
            "%(asctime)s | %(levelname)s | %(name)s | %(message)s"
        )
        console = logging.StreamHandler()
        console.setFormatter(formato)
        arquivo = logging.FileHandler(caminho_log, encoding="utf-8")
        arquivo.setFormatter(formato)
        LOGGER.addHandler(console)
        LOGGER.addHandler(arquivo)
        LOGGER.propagate = False
    return caminho_log


def carregar_usuarios(caminho_planilha: str | Path) -> list[dict[str, str]]:
    """Carrega da planilha mestra os registros extraidos com sucesso."""
    caminho = Path(caminho_planilha).expanduser().resolve()
    if not caminho.is_file():
        raise FileNotFoundError(f"Planilha mestra nao encontrada: {caminho}")
    if caminho.suffix.lower() != ".xlsx":
        raise ValueError(f"A planilha deve possuir extensao .xlsx: {caminho}")

    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = workbook[NOME_ABA] if NOME_ABA in workbook.sheetnames else workbook.active
        linhas = planilha.iter_rows(values_only=True)
        cabecalho = next(linhas, None)
        if not cabecalho:
            return []

        colunas = {
            str(nome).strip(): indice
            for indice, nome in enumerate(cabecalho)
            if nome is not None
        }
        obrigatorias = {"Nome", "CPF", "E-mail", "Telefone", "Nascimento", "Endereco"}
        ausentes = sorted(obrigatorias - colunas.keys())
        if ausentes:
            raise ValueError(
                "Colunas obrigatorias ausentes na planilha: " + ", ".join(ausentes)
            )

        usuarios: list[dict[str, str]] = []
        for numero_linha, linha in enumerate(linhas, start=2):
            registro = {
                nome: _texto(linha[indice]) if indice < len(linha) else ""
                for nome, indice in colunas.items()
            }
            if not _registro_aprovado(registro):
                continue

            nome, sobrenome = _separar_nome(registro["Nome"])
            usuarios.append(
                {
                    "nome": nome,
                    "sobrenome": sobrenome,
                    "cpf": _normalizar_cpf(registro["CPF"], numero_linha),
                    "email": registro["E-mail"],
                    "telefone": registro["Telefone"],
                    "nascimento": _formatar_nascimento(
                        linha[colunas["Nascimento"]]
                    ),
                    "endereco": registro["Endereco"],
                    "observacao": "Cadastro importado da planilha mestra",
                    "status": "ATIVO",
                }
            )
        return usuarios
    finally:
        workbook.close()


def _registro_aprovado(registro: dict[str, str]) -> bool:
    sucesso = registro.get("Sucesso", "Sim").strip().casefold()
    return sucesso in {"sim", "true", "1", "sucesso"}


def _separar_nome(nome_completo: str) -> tuple[str, str]:
    partes = (nome_completo or "").strip().split(maxsplit=1)
    if not partes:
        return "Cliente", "Silva"
    if len(partes) == 1:
        return partes[0], partes[0]
    return partes[0], partes[1]


def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def _normalizar_cpf(valor: Any, numero_linha: int | None = None) -> str:
    """Remove a pontuacao e garante os 11 digitos exigidos pelo portal."""
    cpf = re.sub(r"\D", "", _texto(valor))
    if len(cpf) != 11:
        local = f" na linha {numero_linha}" if numero_linha is not None else ""
        raise ValueError(f"CPF invalido{local}: deve conter exatamente 11 digitos.")
    return cpf


def _formatar_nascimento(valor: Any) -> str:
    """Converte a data para AAAA-MM-DD, formato aceito pelo input HTML date."""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")

    texto = _texto(valor)
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return texto


def chamar_servico_json(
    url: str,
    dados: dict[str, Any],
    token: str = "",
    timeout: float = 10.0,
) -> dict[str, Any]:
    """Envia JSON por POST e devolve a resposta JSON do serviço."""
    cabecalhos = {"Content-Type": "application/json", "Accept": "application/json"}
    if token:
        cabecalhos["Authorization"] = f"Bearer {token}"
    requisicao = Request(
        url,
        data=json.dumps(dados, ensure_ascii=False).encode("utf-8"),
        headers=cabecalhos,
        method="POST",
    )

    try:
        with urlopen(requisicao, timeout=timeout) as resposta:
            conteudo = resposta.read().decode("utf-8")
            if not 200 <= resposta.status < 300:
                raise IntegracaoAPIError(
                    f"Servico retornou HTTP {resposta.status}: {url}"
                )
    except HTTPError as erro:
        detalhe = erro.read().decode("utf-8", errors="replace")
        raise IntegracaoAPIError(
            f"Servico retornou HTTP {erro.code}: {detalhe or erro.reason}"
        ) from erro
    except (URLError, TimeoutError, OSError) as erro:
        raise IntegracaoAPIError(f"Falha de conexao com {url}: {erro}") from erro

    if not conteudo.strip():
        return {}
    try:
        retorno = json.loads(conteudo)
    except json.JSONDecodeError as erro:
        raise IntegracaoAPIError("Servico retornou JSON invalido.") from erro
    if not isinstance(retorno, dict):
        raise IntegracaoAPIError("Servico deve retornar um objeto JSON.")
    return retorno


def enriquecer_usuario_ml(
    usuario: dict[str, str], ml_url: str, token: str = "", timeout: float = 10.0
) -> dict[str, str]:
    """Enriquece/valida um cadastro via ML; em falha, preserva os dados originais."""
    if not ml_url:
        return usuario.copy()
    try:
        retorno = chamar_servico_json(ml_url, {"dados": usuario}, token, timeout)
        dados_ml = retorno.get("dados", retorno)
        if not isinstance(dados_ml, dict):
            raise IntegracaoAPIError("Resposta do ML nao contem um objeto 'dados'.")
        enriquecido = usuario.copy()
        enriquecido.update(
            {chave: _texto(valor) for chave, valor in dados_ml.items() if chave in usuario}
        )
        enriquecido["cpf"] = _normalizar_cpf(enriquecido["cpf"])
        LOGGER.info("Dados enriquecidos pelo ML para CPF final %s.", _mascarar_cpf(enriquecido["cpf"]))
        return enriquecido
    except (IntegracaoAPIError, ValueError) as erro:
        LOGGER.warning("ML indisponivel; usando dados originais como fallback: %s", erro)
        return usuario.copy()


def cadastrar_usuarios_api(
    usuarios: list[dict[str, str]],
    api_url: str,
    ml_url: str = "",
    token: str = "",
    timeout: float = 10.0,
) -> tuple[int, list[dict[str, str]]]:
    """Cadastra via API e retorna quantidade de sucessos e itens para fallback."""
    cadastrados = 0
    falhas: list[dict[str, str]] = []
    for usuario in usuarios:
        preparado = enriquecer_usuario_ml(usuario, ml_url, token, timeout)
        try:
            retorno = chamar_servico_json(api_url, preparado, token, timeout)
            if retorno.get("sucesso") is False:
                raise IntegracaoAPIError(_texto(retorno.get("erro", "Cadastro recusado.")))
            cadastrados += 1
            LOGGER.info(
                "Cadastro via API concluido para CPF %s.",
                _mascarar_cpf(preparado["cpf"]),
            )
        except IntegracaoAPIError as erro:
            LOGGER.error(
                "Cadastro via API falhou para CPF %s: %s",
                _mascarar_cpf(preparado["cpf"]),
                erro,
            )
            falhas.append(preparado)
    return cadastrados, falhas


def _mascarar_cpf(cpf: str) -> str:
    digitos = re.sub(r"\D", "", cpf)
    return f"***.***.***-{digitos[-2:]}" if len(digitos) == 11 else "CPF_INVALIDO"


def preencher_portal_rapido(
    page, usuarios: list[dict[str, str]], qtd: int = 10
) -> tuple[int, list[dict[str, str]]]:
    """Preenche o portal, isolando erros para continuar com os demais cadastros."""
    page.on("dialog", lambda dialog: dialog.accept())
    page.click("#btnClearAll")

    lista = usuarios[:qtd] if qtd else usuarios
    total = len(lista)
    print(f"[RPA Preenchimento] Cadastrando {total} usuarios no Portal Fake...")

    cadastrados = 0
    falhas: list[dict[str, str]] = []
    for usuario in lista:
        try:
            page.click("#btnNovo")
            page.fill("#f_nome", usuario.get("nome", ""))
            page.fill("#f_sobrenome", usuario.get("sobrenome", ""))
            cpf = _normalizar_cpf(usuario.get("cpf", ""))
            page.fill("#f_cpf", cpf)
            cpf_preenchido = page.locator("#f_cpf").input_value()
            if cpf_preenchido != cpf:
                raise RuntimeError(
                    f"Falha ao preencher o CPF. Esperado: {cpf}; preenchido: {cpf_preenchido}."
                )
            page.fill("#f_email", usuario.get("email", ""))
            page.fill("#f_telefone", usuario.get("telefone", ""))
            page.fill("#f_nascimento", usuario.get("nascimento", ""))
            page.fill("#f_endereco", usuario.get("endereco", ""))
            page.fill("#f_observacao", usuario.get("observacao", ""))
            page.select_option("#f_status", (usuario.get("status") or "ATIVO").upper())
            page.click("#btnSalvar")
            page.locator("#modal").wait_for(state="hidden", timeout=3000)
            cadastrados += 1
            LOGGER.info("Cadastro RPA concluido para CPF %s.", _mascarar_cpf(cpf))
        except Exception as erro:
            falhas.append(usuario)
            LOGGER.error(
                "Cadastro RPA falhou para CPF %s: %s",
                _mascarar_cpf(usuario.get("cpf", "")),
                erro,
            )
            try:
                if page.locator("#modal").is_visible():
                    page.click("#btnModalClose")
            except Exception:
                LOGGER.debug("Nao foi possivel fechar o modal apos a falha.", exc_info=True)

    print(f"[RPA Preenchimento] {cadastrados}/{total} cadastros inseridos no portal.")
    return cadastrados, falhas


def salvar_evidencia_cadastros(page, pasta_destino: str | Path) -> Path:
    """Salva uma captura de tela completa do portal após os cadastros."""
    pasta = Path(pasta_destino).expanduser().resolve()
    pasta.mkdir(parents=True, exist_ok=True)
    data_hora = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    caminho = pasta / f"cadastros_concluidos_{data_hora}.png"

    page.wait_for_timeout(500)
    page.screenshot(path=str(caminho), full_page=True)
    print(f"[Evidencia] Screenshot dos cadastros salvo em: {caminho}")
    return caminho


def criar_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Cadastra no Portal Fake os dados da planilha mestra."
    )
    parser.add_argument(
        "planilha",
        nargs="?",
        default=str(PLANILHA_PADRAO),
        help="Caminho do arquivo .xlsx (padrao: resources/planilha_mestra.xlsx).",
    )
    parser.add_argument(
        "--quantidade",
        type=int,
        default=0,
        help="Quantidade maxima de cadastros. Use 0 para cadastrar todos.",
    )
    parser.add_argument(
        "--pasta-screenshots",
        default=str(PASTA_SCREENSHOTS),
        help="Pasta onde a evidencia final em PNG sera salva.",
    )
    parser.add_argument(
        "--modo",
        choices=["auto", "api", "rpa"],
        default="auto",
        help="auto tenta API e usa RPA como fallback; api nao usa fallback; rpa usa o navegador.",
    )
    parser.add_argument(
        "--api-url",
        default=os.getenv("CADASTRO_API_URL", ""),
        help="Endpoint HTTP para cadastro (ou CADASTRO_API_URL no .env).",
    )
    parser.add_argument(
        "--ml-url",
        default=os.getenv("CADASTRO_ML_URL", ""),
        help="Endpoint opcional para enriquecimento ML (ou CADASTRO_ML_URL).",
    )
    parser.add_argument(
        "--timeout-api",
        type=float,
        default=float(os.getenv("CADASTRO_API_TIMEOUT", "10")),
        help="Timeout das integracoes HTTP em segundos.",
    )
    return parser


def executar_rpa(
    usuarios: list[dict[str, str]], pasta_screenshots: str | Path
) -> tuple[int, list[dict[str, str]], Path]:
    """Executa o cadastro via navegador e salva a evidência final."""
    if not INDEX_HTML.is_file():
        raise FileNotFoundError(f"Portal Fake nao encontrado: {INDEX_HTML}")

    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            user_data_dir=str(BROWSER_DATA_DIR),
            headless=False,
        )
        try:
            page = context.new_page()
            page.goto(INDEX_HTML.resolve().as_uri())
            cadastrados, falhas = preencher_portal_rapido(page, usuarios, qtd=0)
            evidencia = salvar_evidencia_cadastros(page, pasta_screenshots)
            return cadastrados, falhas, evidencia
        finally:
            context.close()


def main() -> int:
    argumentos = criar_parser().parse_args()
    caminho_log = configurar_logs()
    LOGGER.info("Inicio do Processo 3. Modo solicitado: %s.", argumentos.modo)

    try:
        usuarios = carregar_usuarios(argumentos.planilha)
        if argumentos.quantidade:
            usuarios = usuarios[: argumentos.quantidade]
        if not usuarios:
            LOGGER.warning("Nenhum registro com sucesso encontrado na planilha mestra.")
            return 1

        token = os.getenv("CADASTRO_API_TOKEN", "")
        cadastrados_api = 0
        cadastrados_rpa = 0
        falhas = usuarios

        if argumentos.modo in {"auto", "api"}:
            if argumentos.api_url:
                cadastrados_api, falhas = cadastrar_usuarios_api(
                    usuarios,
                    api_url=argumentos.api_url,
                    ml_url=argumentos.ml_url,
                    token=token,
                    timeout=argumentos.timeout_api,
                )
            elif argumentos.modo == "api":
                raise IntegracaoAPIError(
                    "CADASTRO_API_URL nao configurada para o modo api."
                )
            else:
                LOGGER.warning("API nao configurada; acionando fallback RPA.")

        if argumentos.modo == "rpa":
            falhas = [
                enriquecer_usuario_ml(
                    usuario, argumentos.ml_url, token, argumentos.timeout_api
                )
                for usuario in usuarios
            ]

        if argumentos.modo == "rpa" or (argumentos.modo == "auto" and falhas):
            LOGGER.warning("Fallback RPA acionado para %s cadastro(s).", len(falhas))
            cadastrados_rpa, falhas, evidencia = executar_rpa(
                falhas, argumentos.pasta_screenshots
            )
            LOGGER.info("Evidencia final salva em %s.", evidencia)

        total_sucesso = cadastrados_api + cadastrados_rpa
        LOGGER.info(
            "Processo finalizado: %s sucesso(s), %s falha(s), API=%s, RPA=%s.",
            total_sucesso,
            len(falhas),
            cadastrados_api,
            cadastrados_rpa,
        )
        print(f"Log da execucao: {caminho_log}")
        return 0 if total_sucesso and not falhas else 1
    except (FileNotFoundError, ValueError, IntegracaoAPIError) as erro:
        LOGGER.error("Processo 3 interrompido: %s", erro)
        return 2
    except Exception:
        LOGGER.exception("Falha inesperada no Processo 3.")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
